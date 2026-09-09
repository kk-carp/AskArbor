from pathlib import Path

import pytest
from docx import Document as DocxDocument
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter

from backend.infra import parsers


def test_parse_document_reads_markdown_utf8(tmp_path: Path) -> None:
    file_path = tmp_path / "course.md"
    file_path.write_text("# 标题\n课程说明", encoding="utf-8")

    assert parsers.parse_document(file_path, "md") == "# 标题\n课程说明"


def test_parse_document_reads_txt_utf8(tmp_path: Path) -> None:
    file_path = tmp_path / "notes.txt"
    file_path.write_text("hello\nfde", encoding="utf-8")

    assert parsers.parse_document(file_path, ".txt") == "hello\nfde"


def test_parse_document_reads_docx(tmp_path: Path) -> None:
    file_path = tmp_path / "guide.docx"
    document = DocxDocument()
    document.add_paragraph("line 1")
    document.add_paragraph("line 2")
    document.save(file_path)

    assert parsers.parse_document(file_path, "docx") == "line 1\nline 2"


def test_parse_document_reads_pdf_text_via_reader_patch(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    file_path = tmp_path / "handbook.pdf"
    file_path.write_bytes(b"%PDF-1.4\n%%EOF")

    class _FakePage:
        @staticmethod
        def extract_text() -> str:
            return "pdf content"

    class _FakeReader:
        def __init__(self, _path: str) -> None:
            self.pages = [_FakePage()]

    monkeypatch.setattr(parsers, "PdfReader", _FakeReader)
    assert parsers.parse_document(file_path, "pdf") == "pdf content"


def test_parse_document_raises_for_scanned_pdf(tmp_path: Path) -> None:
    file_path = tmp_path / "scanned.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    with file_path.open("wb") as output:
        writer.write(output)

    with pytest.raises(ValueError, match="Unsupported scanned PDF"):
        parsers.parse_document(file_path, "pdf")


def test_parse_document_raises_for_empty_text(tmp_path: Path) -> None:
    file_path = tmp_path / "empty.txt"
    file_path.write_text("   \n", encoding="utf-8")

    with pytest.raises(ValueError, match="Document text is empty"):
        parsers.parse_document(file_path, "txt")


def test_parse_document_raises_for_unsupported_extension(tmp_path: Path) -> None:
    file_path = tmp_path / "data.csv"
    file_path.write_text("a,b,c", encoding="utf-8")

    with pytest.raises(ValueError, match="Unsupported file extension"):
        parsers.parse_document(file_path, "csv")


def _blank_layout(presentation: Presentation):
    return presentation.slide_layouts[min(6, len(presentation.slide_layouts) - 1)]


def test_parse_document_reads_pptx_title_body_table_and_notes(tmp_path: Path) -> None:
    file_path = tmp_path / "course.pptx"
    presentation = Presentation()
    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_slide.shapes.title.text = "课程标题"
    if len(title_slide.placeholders) > 1:
        title_slide.placeholders[1].text = "课程正文"
    title_slide.notes_slide.notes_text_frame.text = "讲师备注"

    table_slide = presentation.slides.add_slide(_blank_layout(presentation))
    table = table_slide.shapes.add_table(
        2, 2, Inches(0.5), Inches(0.5), Inches(4), Inches(1.5)
    ).table
    table.cell(0, 0).text = "姓名"
    table.cell(0, 1).text = "学号"
    table.cell(1, 0).text = "张三"
    table.cell(1, 1).text = "001"
    presentation.save(file_path)

    text = parsers.parse_document(file_path, "pptx")
    assert "第 1 页" in text
    assert "课程标题" in text
    assert "课程正文" in text
    assert "讲师备注" in text
    assert "第 2 页" in text
    assert "姓名" in text
    assert "学号" in text
    assert "张三" in text
    assert "001" in text


def test_parse_document_raises_for_empty_pptx(tmp_path: Path) -> None:
    file_path = tmp_path / "empty.pptx"
    Presentation().save(file_path)

    with pytest.raises(ValueError, match="Document text is empty"):
        parsers.parse_document(file_path, "pptx")


def test_parse_document_rejects_legacy_ppt(tmp_path: Path) -> None:
    file_path = tmp_path / "old.ppt"
    file_path.write_bytes(b"legacy-ppt")

    with pytest.raises(ValueError, match="Unsupported file extension"):
        parsers.parse_document(file_path, "ppt")
