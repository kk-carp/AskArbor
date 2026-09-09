"""按格式把文档转成纯文本（md/txt/pdf/docx/pptx）；空文件或无文字则失败。不切片、不访问数据库。"""

from pathlib import Path

from docx import Document as DocxDocument
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader


def _normalize_extension(extension: str) -> str:
    normalized = extension.lower().strip()
    if normalized.startswith("."):
        normalized = normalized[1:]
    return normalized


def _ensure_non_empty_text(text: str) -> str:
    if not text.strip():
        raise ValueError("Document text is empty")
    return text


def _parse_markdown_or_text(path: Path) -> str:
    content = path.read_text(encoding="utf-8")
    return _ensure_non_empty_text(content)


def _parse_docx(path: Path) -> str:
    document = DocxDocument(path)
    lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    return _ensure_non_empty_text("\n".join(lines))


def _parse_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    pages: list[str] = []
    for page in reader.pages:
        extracted = page.extract_text() or ""
        if extracted.strip():
            pages.append(extracted)
    if not pages:
        raise ValueError("Unsupported scanned PDF")
    return _ensure_non_empty_text("\n".join(pages))


def _shape_lines(shape: object) -> list[str]:
    """递归抽取文本框与表格；不读图表、SmartArt、嵌入图片。"""
    lines: list[str] = []
    if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
        for child in shape.shapes:
            lines.extend(_shape_lines(child))
        return lines
    if getattr(shape, "has_text_frame", False):
        for paragraph in shape.text_frame.paragraphs:
            text = (paragraph.text or "").strip()
            if text:
                lines.append(text)
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            cells = [(cell.text or "").strip() for cell in row.cells]
            cells = [cell for cell in cells if cell]
            if cells:
                lines.append("\t".join(cells))
    return lines


def _parse_pptx(path: Path) -> str:
    presentation = Presentation(str(path))
    pages: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        lines: list[str] = []
        for shape in slide.shapes:
            lines.extend(_shape_lines(shape))
        if slide.has_notes_slide:
            notes = (slide.notes_slide.notes_text_frame.text or "").strip()
            if notes:
                lines.append(notes)
        if lines:
            pages.append(f"第 {index} 页\n" + "\n".join(lines))
    if not pages:
        raise ValueError("Document text is empty")
    return _ensure_non_empty_text("\n\n".join(pages))


def parse_document(path: Path, extension: str) -> str:
    """将支持格式文档解析为纯文本。"""
    normalized = _normalize_extension(extension)
    if normalized in {"md", "txt"}:
        return _parse_markdown_or_text(path)
    if normalized == "docx":
        return _parse_docx(path)
    if normalized == "pdf":
        return _parse_pdf(path)
    if normalized == "pptx":
        return _parse_pptx(path)
    raise ValueError(f"Unsupported file extension: {normalized}")
